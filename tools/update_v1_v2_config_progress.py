"""Expand the V1 classification workbook into a UTF-8 V2 progress CSV.

The historical .csv file is an XLSX container. Read it without changing the
source, and preserve manually reviewed item statuses on subsequent runs.
"""

import csv
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "V1_CONFIG_BY_SEGMENT_20260925.csv"
SUMMARY = ROOT / "docs" / "V1_V2_SUBCATEGORY_PROGRESS_20260926.csv"
OUTPUT = ROOT / "docs" / "V1_V2_CONFIG_PROGRESS_20260926.csv"
REVIEW_FIELDS = ("V2配置项进度", "V2实现依据", "下一步")


def read_csv(path):
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def main():
    summary_rows = read_csv(SUMMARY)
    summaries = {row["子类"]: row for row in summary_rows}
    if len(summaries) != 23 or len(summary_rows) != 23:
        raise ValueError("expected 23 distinct V1 subcategories")

    with SOURCE.open("rb") as handle:
        sheet = load_workbook(handle, read_only=True, data_only=True).active
        rows = sheet.iter_rows(values_only=True)
        source_fields = next(rows)
        source_rows = [dict(zip(source_fields, values)) for values in rows]
    if len(source_rows) != 1010:
        raise ValueError(f"expected 1010 V1 items, got {len(source_rows)}")

    counts = Counter(row["子类"] for row in source_rows)
    for code, summary in summaries.items():
        if counts[code] != int(summary["V1条数"]):
            raise ValueError(f"count mismatch for {code}: {counts[code]}")
    if len({row["配置键"] for row in source_rows}) != 1010:
        raise ValueError("V1 config keys are not unique")

    existing = {row["配置键"]: row for row in read_csv(OUTPUT)}
    fields = list(source_fields) + ["V2段页", "V2专题页", "V2子类进度", *REVIEW_FIELDS]
    with OUTPUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for source in source_rows:
            summary = summaries[source["子类"]]
            prior = existing.get(source["配置键"], {})
            row = {**source, "V2段页": summary["V2段页"],
                   "V2专题页": summary["V2专题页"],
                   "V2子类进度": summary["子类进度"],
                   "V2配置项进度": prior.get("V2配置项进度") or "待逐项核对",
                   "V2实现依据": prior.get("V2实现依据") or "",
                   "下一步": prior.get("下一步") or ""}
            writer.writerow(row)
    print(f"wrote {len(source_rows)} V1 items across {len(summaries)} subcategories to {OUTPUT}")


if __name__ == "__main__":
    main()
