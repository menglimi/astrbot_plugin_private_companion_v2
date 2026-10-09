"""Read document text with standard parsers without executing attachment content."""

import hashlib
import io
from pathlib import PurePosixPath
from zipfile import ZipFile, BadZipFile

from companion_v2.adapters.astrbot.files import DOCUMENT_SUFFIXES, TEXT_SUFFIXES, FileText
from companion_v2.execution.models import TaskError


def _archive(content):
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            items = archive.infolist()
            if (len(items) > 3000 or sum(item.file_size for item in items) > 64 * 1024 * 1024
                    or any(item.flag_bits & 1 for item in items)):
                raise TaskError("host_file_limit", "document archive exceeds its extraction budget")
    except BadZipFile:
        raise TaskError("host_file_invalid", "document archive is invalid") from None


def _text_parts(content, suffix):
    if suffix in TEXT_SUFFIXES:
        encoding = "utf-16" if content.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            text = content.decode("gb18030")
        if "\x00" in text:
            raise TaskError("host_file_invalid", "attachment is not a text document")
        yield text
    elif suffix == ".pdf":
        import fitz

        with fitz.open(stream=content, filetype="pdf") as document:
            if document.needs_pass:
                raise TaskError("host_file_empty", "encrypted document has no authorized readable text")
            for index, page in enumerate(document):
                yield f"\n[第 {index + 1} 页]\n" + page.get_text(sort=True) if page.get_text().strip() else ""
    elif suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        for paragraph in document.paragraphs:
            yield paragraph.text + "\n"
        for table in document.tables:
            for row in table.rows:
                yield " | ".join(cell.text for cell in row.cells) + "\n"
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook:
                yield f"\n[工作表 {sheet.title}]\n"
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        yield " | ".join(values) + "\n"
        finally:
            workbook.close()
    elif suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(content))
        for index, slide in enumerate(presentation.slides):
            yield f"\n[幻灯片 {index + 1}]\n"
            for shape in slide.shapes:
                if shape.has_text_frame:
                    yield shape.text + "\n"
                if shape.has_table:
                    for row in shape.table.rows:
                        yield " | ".join(cell.text for cell in row.cells) + "\n"


def _allowed_suffixes(formats):
    if formats == "text_only":
        return TEXT_SUFFIXES
    if formats == "text_and_documents":
        return TEXT_SUFFIXES | DOCUMENT_SUFFIXES
    if isinstance(formats, (set, frozenset, tuple)):
        selected = frozenset(formats)
        if all(type(item) is str for item in selected) and selected <= (TEXT_SUFFIXES | DOCUMENT_SUFFIXES):
            return selected
    raise TaskError("host_file_format_unsupported", "this document format selection is invalid")


def extract_document(content, position, name, *, maximum_chars=24000, formats="text_and_documents"):
    suffix = PurePosixPath(name).suffix.lower()
    allowed = _allowed_suffixes(formats)
    if suffix not in allowed:
        raise TaskError("host_file_format_unsupported", "this document format is not selected or supported")
    if type(content) is not bytes or not 0 < len(content) <= 16 * 1024 * 1024:
        raise TaskError("host_file_limit", "document exceeds its original byte budget")
    if suffix in (".docx", ".xlsx", ".pptx"):
        _archive(content)
    pieces, size, truncated = [], 0, False
    try:
        for piece in _text_parts(content, suffix):
            remaining = maximum_chars - size
            if len(piece) > remaining:
                pieces.append(piece[:remaining])
                truncated = True
                break
            pieces.append(piece)
            size += len(piece)
    except ImportError:
        raise TaskError("host_file_decoder_unavailable", "document parser is unavailable") from None
    except TaskError:
        raise
    except Exception:
        raise TaskError("host_file_invalid", "document text could not be extracted") from None
    text = "".join(pieces).strip()
    if not text:
        raise TaskError("host_file_empty", "document contains no readable text")
    return FileText(position, name, hashlib.sha256(content).hexdigest(), len(content), text, truncated)
